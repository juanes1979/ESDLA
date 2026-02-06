"""
Script to update the equipment catalog with ALL categories from utumno.xlsm
Including:
- Weapons (categorized: Sencillas C/C, Sencillas Distancia, Marciales C/C, Marciales Distancia)
- Armors (categorized: Ligeras, Medias, Pesadas, Escudos)
- Tools (Herramientas)
- Games (Juegos)
- Musical Instruments (Instrumentos)
- General Equipment
- Food/Consumables (Consumibles)
- Inn/Restaurant Food (Comida en Posadas)
- Medicinal Herbs/Potions (Hierbas)
- Poisons (Venenos)
- Mounts (Monturas)
- Mount Accessories (Accesorios de Monturas)
- Land Transport (Transporte Terrestre)
- Sea Transport (Transporte Marítimo)
- Construction Materials (Elementos de Construcción)
"""
import openpyxl
import asyncio
from motor.motor_asyncio import AsyncIOMotorClient
import os
from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

def safe_str(val, max_len=500):
    """Safely convert value to string"""
    if val is None:
        return ""
    return str(val)[:max_len].strip()

def safe_float(val):
    """Safely convert to float"""
    if val is None:
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None

def safe_int(val):
    """Safely convert to int"""
    if val is None:
        return None
    try:
        return int(val)
    except (ValueError, TypeError):
        return None

def parse_currency(uds, mon):
    """Parse price from uds and mon columns"""
    try:
        precio = safe_float(uds) if uds else None
        moneda = safe_str(mon) if mon else "mp"
        return precio, moneda
    except:
        return None, "mp"

async def update_equipment_catalog():
    wb = openpyxl.load_workbook('/app/data/utumno.xlsm', data_only=True)
    ws = wb['Equipo']
    
    catalog = {
        "herramientas": [],
        "juegos": [],
        "instrumentos_musicales": [],
        "equipo_general": [],
        "consumibles": [],
        "comida_posadas": [],
        "hierbas": [],
        "venenos": [],
        "armas_sencillas_cc": [],
        "armas_sencillas_distancia": [],
        "armas_marciales_cc": [],
        "armas_marciales_distancia": [],
        "armaduras_ligeras": [],
        "armaduras_medias": [],
        "armaduras_pesadas": [],
        "escudos": [],
        "monturas": [],
        "accesorios_monturas": [],
        "transporte_terrestre": [],
        "transporte_maritimo": [],
        "construccion": []
    }
    
    # === HERRAMIENTAS (B3-B21, cols B-E) ===
    print("Extracting Herramientas...")
    for row in range(3, 22):
        nombre = safe_str(ws.cell(row=row, column=2).value)
        if not nombre:
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=3).value,
            ws.cell(row=row, column=4).value
        )
        peso = safe_float(ws.cell(row=row, column=5).value)
        catalog["herramientas"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso
        })
    
    # === JUEGOS (B25-B31, cols B-E) ===
    print("Extracting Juegos...")
    for row in range(25, 32):
        nombre = safe_str(ws.cell(row=row, column=2).value)
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=3).value,
            ws.cell(row=row, column=4).value
        )
        peso = safe_float(ws.cell(row=row, column=5).value)
        catalog["juegos"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso
        })
    
    # === INSTRUMENTOS MUSICALES (B34-B43, cols B-E) ===
    print("Extracting Instrumentos Musicales...")
    for row in range(34, 44):
        nombre = safe_str(ws.cell(row=row, column=2).value)
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=3).value,
            ws.cell(row=row, column=4).value
        )
        peso = safe_float(ws.cell(row=row, column=5).value)
        catalog["instrumentos_musicales"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso
        })
    
    # === EQUIPO GENERAL (G3-G111, cols G-J) - excluding food ===
    print("Extracting Equipo General...")
    for row in range(3, 112):
        nombre = safe_str(ws.cell(row=row, column=7).value)
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=8).value,
            ws.cell(row=row, column=9).value
        )
        peso = safe_float(ws.cell(row=row, column=10).value)
        catalog["equipo_general"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso
        })
    
    # === CONSUMIBLES Y ALIMENTACIÓN (G114-G184, cols G-J) ===
    print("Extracting Consumibles...")
    for row in range(114, 185):
        nombre = safe_str(ws.cell(row=row, column=7).value)
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=8).value,
            ws.cell(row=row, column=9).value
        )
        peso = safe_float(ws.cell(row=row, column=10).value)
        catalog["consumibles"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso
        })
    
    # === COMIDA EN POSADAS (G188-G223, cols G-J) ===
    print("Extracting Comida en Posadas...")
    for row in range(188, 224):
        nombre = safe_str(ws.cell(row=row, column=7).value)
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=8).value,
            ws.cell(row=row, column=9).value
        )
        peso = safe_float(ws.cell(row=row, column=10).value)
        catalog["comida_posadas"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso
        })
    
    # === HIERBAS MEDICINALES / POCIONES (G227-G313, cols G-L) ===
    print("Extracting Hierbas...")
    for row in range(227, 314):
        nombre = safe_str(ws.cell(row=row, column=7).value)
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=8).value,
            ws.cell(row=row, column=9).value
        )
        peso = safe_float(ws.cell(row=row, column=10).value)
        forma = safe_str(ws.cell(row=row, column=11).value)
        efecto = safe_str(ws.cell(row=row, column=12).value)
        catalog["hierbas"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso,
            "forma_preparacion": forma,
            "efecto": efecto
        })
    
    # === VENENOS (G317-G331, cols G-L) ===
    print("Extracting Venenos...")
    for row in range(317, 332):
        nombre = safe_str(ws.cell(row=row, column=7).value)
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=8).value,
            ws.cell(row=row, column=9).value
        )
        peso = safe_float(ws.cell(row=row, column=10).value)
        forma = safe_str(ws.cell(row=row, column=11).value)
        efecto = safe_str(ws.cell(row=row, column=12).value)
        catalog["venenos"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso,
            "forma_preparacion": forma,
            "efecto": efecto
        })
    
    # === ARMAS (N3-V35) ===
    # Sencillas C/C: rows 4-11
    # Sencillas Distancia: rows 13-17
    # Marciales C/C: rows 19-32
    # Marciales Distancia: rows 34-35
    print("Extracting Armas...")
    
    def extract_weapon(row):
        nombre = safe_str(ws.cell(row=row, column=14).value)  # N
        if not nombre or nombre.lower().startswith('arma'):
            return None
        precio, moneda = parse_currency(
            ws.cell(row=row, column=15).value,  # O
            ws.cell(row=row, column=16).value   # P
        )
        peso = safe_float(ws.cell(row=row, column=17).value)  # Q
        modificador = safe_str(ws.cell(row=row, column=18).value)  # R
        dano = safe_str(ws.cell(row=row, column=19).value)  # S
        herida = safe_int(ws.cell(row=row, column=20).value)  # T
        distancia = safe_str(ws.cell(row=row, column=21).value)  # U
        return {
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso,
            "modificador": modificador,
            "dano": dano,
            "herida": herida,
            "alcance": distancia
        }
    
    # Armas sencillas cuerpo a cuerpo (rows 4-11)
    for row in range(4, 12):
        weapon = extract_weapon(row)
        if weapon:
            catalog["armas_sencillas_cc"].append(weapon)
    
    # Armas sencillas a distancia (rows 13-17)
    for row in range(13, 18):
        weapon = extract_weapon(row)
        if weapon:
            catalog["armas_sencillas_distancia"].append(weapon)
    
    # Armas marciales cuerpo a cuerpo (rows 19-32)
    for row in range(19, 33):
        weapon = extract_weapon(row)
        if weapon:
            catalog["armas_marciales_cc"].append(weapon)
    
    # Armas marciales a distancia (rows 34-35)
    for row in range(34, 36):
        weapon = extract_weapon(row)
        if weapon:
            catalog["armas_marciales_distancia"].append(weapon)
    
    # === ARMADURAS (W3-AD21) ===
    print("Extracting Armaduras...")
    
    def extract_armor(row):
        nombre = safe_str(ws.cell(row=row, column=23).value)  # W
        if not nombre or nombre.lower().startswith('armadura') or nombre.lower() == 'escudos':
            return None
        precio, moneda = parse_currency(
            ws.cell(row=row, column=24).value,  # X
            ws.cell(row=row, column=25).value   # Y
        )
        peso = safe_float(ws.cell(row=row, column=26).value)  # Z
        comentarios = safe_str(ws.cell(row=row, column=27).value)  # AA
        ca = safe_int(ws.cell(row=row, column=28).value)  # AB
        otros = safe_str(ws.cell(row=row, column=29).value)  # AC
        return {
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso,
            "ca": ca,
            "comentarios": comentarios,
            "otros": otros
        }
    
    # Armaduras ligeras (rows 4-5)
    for row in range(4, 6):
        armor = extract_armor(row)
        if armor:
            catalog["armaduras_ligeras"].append(armor)
    
    # Armaduras medias (rows 9-11)
    for row in range(9, 12):
        armor = extract_armor(row)
        if armor:
            catalog["armaduras_medias"].append(armor)
    
    # Armaduras pesadas (rows 15-17)
    for row in range(15, 18):
        armor = extract_armor(row)
        if armor:
            catalog["armaduras_pesadas"].append(armor)
    
    # Escudos (row 21)
    armor = extract_armor(21)
    if armor:
        catalog["escudos"].append(armor)
    
    # === MONTURAS (AI3-AI10) ===
    print("Extracting Monturas...")
    for row in range(3, 11):
        nombre = safe_str(ws.cell(row=row, column=35).value)  # AI
        if not nombre or nombre.lower() == 'montura':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=36).value,  # AJ
            ws.cell(row=row, column=37).value   # AK
        )
        carga = safe_int(ws.cell(row=row, column=38).value)  # AL
        constitucion = safe_str(ws.cell(row=row, column=39).value)  # AM
        velocidad = safe_int(ws.cell(row=row, column=40).value)  # AN
        capacidad = safe_str(ws.cell(row=row, column=41).value)  # AO
        catalog["monturas"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "capacidad_carga": carga,
            "constitucion": constitucion,
            "velocidad": velocidad,
            "capacidad_monta": capacidad
        })
    
    # === ACCESORIOS DE MONTURAS (AI14-AI22) ===
    print("Extracting Accesorios de Monturas...")
    for row in range(14, 23):
        nombre = safe_str(ws.cell(row=row, column=35).value)  # AI
        if not nombre or nombre.lower() == 'objeto':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=36).value,  # AJ
            ws.cell(row=row, column=37).value   # AK
        )
        peso = safe_str(ws.cell(row=row, column=38).value)  # AL (can be "—")
        peso_val = None
        if peso and peso != '—':
            try:
                peso_val = float(peso.replace(' kg', '').replace('kg', ''))
            except:
                pass
        catalog["accesorios_monturas"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso_val
        })
    
    # === TRANSPORTE TERRESTRE (AI26-AI35) ===
    print("Extracting Transporte Terrestre...")
    for row in range(26, 36):
        nombre = safe_str(ws.cell(row=row, column=35).value)  # AI
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=36).value,  # AJ
            ws.cell(row=row, column=37).value   # AK
        )
        capacidad = safe_int(ws.cell(row=row, column=38).value)  # AL
        catalog["transporte_terrestre"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "capacidad_kg": capacidad
        })
    
    # === TRANSPORTE MARÍTIMO (AI39-AI44+) ===
    print("Extracting Transporte Marítimo...")
    for row in range(39, 50):
        nombre = safe_str(ws.cell(row=row, column=35).value)  # AI
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=36).value,  # AJ
            ws.cell(row=row, column=37).value   # AK
        )
        capacidad = safe_int(ws.cell(row=row, column=38).value)  # AL
        catalog["transporte_maritimo"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "capacidad_kg": capacidad
        })
    
    # === ELEMENTOS DE CONSTRUCCIÓN (AR3-AV54) ===
    print("Extracting Elementos de Construcción...")
    for row in range(3, 55):
        nombre = safe_str(ws.cell(row=row, column=44).value)  # AR
        if not nombre or nombre.lower() == 'equipo':
            continue
        precio, moneda = parse_currency(
            ws.cell(row=row, column=45).value,  # AS
            ws.cell(row=row, column=46).value   # AT
        )
        peso = safe_float(ws.cell(row=row, column=47).value)  # AU
        m2 = safe_str(ws.cell(row=row, column=48).value)  # AV
        catalog["construccion"].append({
            "nombre": nombre,
            "precio": precio,
            "moneda": moneda,
            "peso_kg": peso,
            "m2": m2
        })
    
    # Update equipment_catalog in MongoDB
    print("\nUpdating MongoDB...")
    await db.equipment_catalog.update_one(
        {"_id": "main"},
        {"$set": catalog},
        upsert=True
    )
    
    # Also update equipment_lists with games and instruments
    await db.equipment_lists.update_one(
        {"_id": "main"},
        {"$set": {
            "juegos": [item["nombre"] for item in catalog["juegos"]],
            "instrumentos_musicales": [item["nombre"] for item in catalog["instrumentos_musicales"]]
        }},
        upsert=True
    )
    
    print("\n=== Summary ===")
    for key, items in catalog.items():
        print(f"{key}: {len(items)} items")
    
    print("\nDone!")

if __name__ == "__main__":
    asyncio.run(update_equipment_catalog())
