"""
Data Extractor for LOTR 5e RPG
Extracts all game data from Excel files and prepares for MongoDB seeding
"""
import openpyxl
import json
from pathlib import Path
from typing import Dict, List, Any, Optional

DATA_DIR = Path(__file__).parent.parent / 'data'

def clean_value(val: Any) -> Any:
    """Clean and normalize cell values"""
    if val is None:
        return None
    if isinstance(val, str):
        val = val.strip()
        if val == '' or val == '-':
            return None
    return val

def extract_cultures(wb: openpyxl.Workbook) -> List[Dict]:
    """Extract cultures from Culturas sheet"""
    sheet = wb['Culturas']
    cultures = []
    
    for row in range(2, sheet.max_row + 1):
        name = clean_value(sheet.cell(row=row, column=1).value)
        if not name:
            continue
            
        culture = {
            "nombre": name,
            "descripcion": clean_value(sheet.cell(row=row, column=2).value),
            "edad_min": clean_value(sheet.cell(row=row, column=3).value),
            "edad_max": clean_value(sheet.cell(row=row, column=4).value),
            "altura_min": clean_value(sheet.cell(row=row, column=5).value),
            "altura_max": clean_value(sheet.cell(row=row, column=6).value),
            "mod_peso_porcentaje": clean_value(sheet.cell(row=row, column=7).value),
            "velocidad": clean_value(sheet.cell(row=row, column=8).value),
            "descanso_horas": clean_value(sheet.cell(row=row, column=9).value),
            "tamano": clean_value(sheet.cell(row=row, column=10).value),
            "nivel_vida": clean_value(sheet.cell(row=row, column=11).value),
            "descripcion_riqueza": clean_value(sheet.cell(row=row, column=12).value),
        }
        cultures.append(culture)
    
    return cultures

def extract_characteristics_by_culture(wb: openpyxl.Workbook) -> Dict[str, Dict]:
    """Extract characteristic modifiers per culture from Características sheet"""
    sheet = wb['Características']
    
    # Get culture names from row 2
    culture_names = []
    for col in range(2, sheet.max_column + 1):
        name = clean_value(sheet.cell(row=2, column=col).value)
        if name:
            culture_names.append((col, name))
    
    # Get characteristic names (FUERZA, DESTREZA, etc.) from column 1
    char_rows = {}
    for row in range(3, 9):  # Rows 3-8 contain FUE, DES, CON, INT, SAB, CAR
        char_name = clean_value(sheet.cell(row=row, column=1).value)
        if char_name:
            char_rows[row] = char_name
    
    # Build culture -> characteristics mapping
    culture_characteristics = {}
    for col, culture_name in culture_names:
        characteristics = {}
        for row, char_name in char_rows.items():
            val = clean_value(sheet.cell(row=row, column=col).value)
            if val is not None:
                try:
                    characteristics[char_name.lower()] = int(val)
                except (ValueError, TypeError):
                    characteristics[char_name.lower()] = val
        
        if characteristics:
            culture_characteristics[culture_name] = characteristics
    
    return culture_characteristics

def extract_backgrounds(wb: openpyxl.Workbook) -> List[Dict]:
    """Extract backgrounds from Trasfondo sheet - complex structure"""
    sheet = wb['Trasfondo']
    backgrounds = []
    
    # This sheet has a complex structure with cultures in columns
    # Row 2: Culture groups (ELFOS, ENANOS, etc.)
    # Row 3: Subcultures
    # Row 4: Background names
    # Row 5: Description
    # Rows 6-7: Skill competencies
    # Row 11: Tool competencies
    # Row 16-17: Rasgos (traits)
    
    current_col = 2
    while current_col <= sheet.max_column:
        subculture = clean_value(sheet.cell(row=3, column=current_col).value)
        bg_name = clean_value(sheet.cell(row=4, column=current_col).value)
        
        if not bg_name:
            current_col += 1
            continue
        
        # Extract traits from rows 16-17
        rasgos = []
        for row in [16, 17]:
            rasgo = clean_value(sheet.cell(row=row, column=current_col).value)
            if rasgo and rasgo not in rasgos:
                rasgos.append(rasgo)
        
        background = {
            "nombre": bg_name,
            "cultura": subculture,
            "descripcion": clean_value(sheet.cell(row=5, column=current_col).value),
            "competencias_habilidades": [],
            "competencias_herramientas": [],
            "idiomas": [],
            "equipo_inicial": [],
            "rasgos": rasgos,  # Lista de rasgos del trasfondo
            "rasgo_distintivo": rasgos[0] if rasgos else None,  # Primer rasgo como distintivo
        }
        
        # Extract skill competencies (rows 6-7 typically)
        for row in range(6, 9):
            skill = clean_value(sheet.cell(row=row, column=current_col).value)
            if skill and not skill.startswith(('Herramienta', 'Idioma', 'Equipo', 'herramienta')):
                background["competencias_habilidades"].append(skill)
        
        # Extract tool competencies (row 11)
        tool = clean_value(sheet.cell(row=11, column=current_col).value)
        if tool:
            background["competencias_herramientas"].append(tool)
        
        backgrounds.append(background)
        current_col += 1
    
    return backgrounds

def extract_occupations(wb: openpyxl.Workbook) -> List[Dict]:
    """Extract occupations/classes from Ocupaciones sheet"""
    sheet = wb['Ocupaciones']
    occupations = []
    
    # Columns B-G contain the 6 main occupations
    occupation_cols = {
        2: "Explorador",
        3: "Guerrero", 
        4: "Lider",
        5: "Maestro",
        6: "Protector",
        7: "Trotamundos"
    }
    
    for col, tipo in occupation_cols.items():
        # Row 2: Vocación name
        vocacion = clean_value(sheet.cell(row=2, column=col).value)
        
        occupation = {
            "tipo": tipo,
            "vocacion": vocacion,
            "dado_golpe": clean_value(sheet.cell(row=3, column=col).value),
            "puntos_golpe_nivel1": clean_value(sheet.cell(row=4, column=col).value),
            "caracteristicas_principales": [],
            "competencia_tiradas_salvacion": [],
            "habilidades_disponibles": [],
            "competencia_armaduras": [],
            "competencia_armas": [],
            "rasgos_por_nivel": {},
        }
        
        # Características principales (rows 5-6)
        for row in range(5, 7):
            char = clean_value(sheet.cell(row=row, column=col).value)
            if char:
                occupation["caracteristicas_principales"].append(char)
        
        # Tiradas de salvación (rows 7-8)
        for row in range(7, 9):
            char = clean_value(sheet.cell(row=row, column=col).value)
            if char:
                occupation["competencia_tiradas_salvacion"].append(char)
        
        occupations.append(occupation)
    
    return occupations

def extract_virtues(wb: openpyxl.Workbook) -> List[Dict]:
    """Extract virtues from Virtudes sheet"""
    sheet = wb['Virtudes']
    virtues = []
    
    # Virtues are organized by culture columns
    # Row 2: Virtue name
    # Row 3: Description
    # Row 4: Traits to mark on character sheet
    # Rows 5-10: Ability score increases
    
    # Start with common virtues (column C = 3)
    for col in range(3, sheet.max_column + 1):
        name = clean_value(sheet.cell(row=2, column=col).value)
        if not name or name == 'VIRTUD':
            continue
        
        # Get culture from header row
        culture_header = clean_value(sheet.cell(row=1, column=col).value)
        
        virtue = {
            "nombre": name,
            "cultura": culture_header if culture_header != "COMUNES" else None,
            "es_comun": culture_header == "COMUNES",
            "descripcion": clean_value(sheet.cell(row=3, column=col).value),
            "rasgos_hoja_pj": clean_value(sheet.cell(row=4, column=col).value),
            "aumentos_caracteristica": {},
        }
        
        # Check ability increases (rows 5-10)
        abilities = ["FUERZA", "DESTREZA", "CONSTITUCIÓN", "INTELIGENCIA", "SABIDURÍA", "CARISMA"]
        for idx, ability in enumerate(abilities, start=5):
            val = clean_value(sheet.cell(row=idx, column=col).value)
            if val and val != '-':
                try:
                    virtue["aumentos_caracteristica"][ability.lower()] = int(val) if val != 'X' else 1
                except:
                    virtue["aumentos_caracteristica"][ability.lower()] = 1
        
        virtues.append(virtue)
    
    return virtues

def extract_equipment(wb: openpyxl.Workbook) -> Dict[str, List[Dict]]:
    """Extract all equipment from Equipo sheet"""
    sheet = wb['Equipo']
    
    equipment = {
        "herramientas": [],
        "equipo_general": [],
        "armas": [],
        "armaduras": [],
    }
    
    # Parse tools (column A-D starting row 3)
    for row in range(3, 50):
        name = clean_value(sheet.cell(row=row, column=1).value)
        if not name or name == 'Equipo':
            continue
        
        tool = {
            "nombre": name,
            "precio": clean_value(sheet.cell(row=row, column=2).value),
            "moneda": clean_value(sheet.cell(row=row, column=3).value),
            "peso_kg": clean_value(sheet.cell(row=row, column=4).value),
        }
        equipment["herramientas"].append(tool)
    
    # Parse general equipment (column E-H starting row 3)
    for row in range(3, 200):
        name = clean_value(sheet.cell(row=row, column=5).value)
        if not name:
            continue
        
        item = {
            "nombre": name,
            "precio": clean_value(sheet.cell(row=row, column=6).value),
            "moneda": clean_value(sheet.cell(row=row, column=7).value),
            "peso_kg": clean_value(sheet.cell(row=row, column=8).value),
        }
        equipment["equipo_general"].append(item)
    
    # Parse weapons (columns starting at N = 14)
    # Structure: Arma(14), Uds(15), Mon(16), Peso(17), Modificador/TipoDano(18), Daño(19), Herida(20), Distancia(21)
    for row in range(3, 100):
        name = clean_value(sheet.cell(row=row, column=14).value)
        if not name or name == 'Arma':
            continue
        
        weapon = {
            "nombre": name,
            "precio": clean_value(sheet.cell(row=row, column=15).value),
            "moneda": clean_value(sheet.cell(row=row, column=16).value),
            "peso_kg": clean_value(sheet.cell(row=row, column=17).value),
            "tipo_dano": clean_value(sheet.cell(row=row, column=18).value),
            "dano": clean_value(sheet.cell(row=row, column=19).value),
            "herida": clean_value(sheet.cell(row=row, column=20).value),
            "alcance": clean_value(sheet.cell(row=row, column=21).value),
        }
        equipment["armas"].append(weapon)
    
    # Parse armors (columns starting at W = 23)
    for row in range(3, 50):
        name = clean_value(sheet.cell(row=row, column=23).value)
        if not name or name == 'Armadura':
            continue
        
        armor = {
            "nombre": name,
            "precio": clean_value(sheet.cell(row=row, column=24).value),
            "moneda": clean_value(sheet.cell(row=row, column=25).value),
            "clase_armadura": clean_value(sheet.cell(row=row, column=26).value),
            "peso_kg": clean_value(sheet.cell(row=row, column=27).value),
            "requisito_fuerza": clean_value(sheet.cell(row=row, column=28).value),
            "desventaja_sigilo": clean_value(sheet.cell(row=row, column=29).value),
        }
        equipment["armaduras"].append(armor)
    
    return equipment

def extract_arts(wb: openpyxl.Workbook) -> List[Dict]:
    """Extract arts/spells from Artes sheet"""
    sheet = wb['Artes']
    arts = []
    
    for row in range(2, sheet.max_row + 1):
        name = clean_value(sheet.cell(row=row, column=1).value)
        if not name:
            continue
        
        art = {
            "nombre": name,
            "resumen_hoja_pj": clean_value(sheet.cell(row=row, column=2).value),
            "descripcion_completa": clean_value(sheet.cell(row=row, column=3).value),
        }
        arts.append(art)
    
    return arts

def extract_patrons(wb: openpyxl.Workbook) -> List[Dict]:
    """Extract patrons from Mecenas sheet"""
    sheet = wb['Mecenas']
    patrons = []
    
    for row in range(2, sheet.max_row + 1):
        name = clean_value(sheet.cell(row=row, column=1).value)
        if not name:
            continue
        
        patron = {
            "nombre": name,
            "ocupaciones_favorecidas": clean_value(sheet.cell(row=row, column=2).value),
            "puntos_comunidad": clean_value(sheet.cell(row=row, column=3).value),
            "ventaja_adicional": clean_value(sheet.cell(row=row, column=4).value),
            "planes": clean_value(sheet.cell(row=row, column=5).value),
        }
        patrons.append(patron)
    
    return patrons

def extract_shadow_rules(wb: openpyxl.Workbook) -> Dict:
    """Extract shadow/corruption rules from Sombra sheet"""
    sheet = wb['Sombra']
    
    shadow = {
        "fuentes_pavor": [],
        "fuentes_avaricia": [],
        "fuentes_desesperacion": [],
        "efectos_sombra": [],
    }
    
    current_section = None
    for row in range(3, sheet.max_row + 1):
        cell_a = clean_value(sheet.cell(row=row, column=1).value)
        
        if cell_a in ['PAVOR', 'AVARICIA', 'DESESPERACIÓN']:
            current_section = cell_a.lower()
            continue
        
        if current_section and cell_a:
            entry = {
                "fuente": cell_a,
                "ejemplo": clean_value(sheet.cell(row=row, column=2).value),
                "puntos_sombra": clean_value(sheet.cell(row=row, column=3).value),
            }
            key = f"fuentes_{current_section}"
            if key in shadow:
                shadow[key].append(entry)
    
    return shadow

def extract_names_by_culture(wb: openpyxl.Workbook) -> Dict[str, Dict]:
    """Extract name generation data from Nombres sheet"""
    sheet = wb['Nombres']
    names_data = {}
    
    # Row 1 has culture names
    current_col = 1
    while current_col <= sheet.max_column:
        culture_name = clean_value(sheet.cell(row=1, column=current_col).value)
        if not culture_name:
            current_col += 1
            continue
        
        # Each culture spans multiple columns for male/female prefixes/suffixes
        names_data[culture_name] = {
            "hombre": {"prefijos": [], "sufijos": []},
            "mujer": {"prefijos": [], "sufijos": []},
            "apellidos": [],
        }
        
        # Extract prefixes/suffixes from rows 4+
        for row in range(4, sheet.max_row + 1):
            prefix_m = clean_value(sheet.cell(row=row, column=current_col).value)
            suffix_m = clean_value(sheet.cell(row=row, column=current_col + 1).value)
            
            if prefix_m:
                names_data[culture_name]["hombre"]["prefijos"].append(prefix_m)
            if suffix_m:
                names_data[culture_name]["hombre"]["sufijos"].append(suffix_m)
        
        current_col += 6  # Move to next culture block
    
    return names_data

def extract_all_data() -> Dict:
    """Extract all game data from Excel files"""
    utumno_path = DATA_DIR / 'utumno.xlsm'
    wb = openpyxl.load_workbook(utumno_path, data_only=True)
    
    data = {
        "culturas": extract_cultures(wb),
        "caracteristicas_por_cultura": extract_characteristics_by_culture(wb),
        "trasfondos": extract_backgrounds(wb),
        "ocupaciones": extract_occupations(wb),
        "virtudes": extract_virtues(wb),
        "equipo": extract_equipment(wb),
        "artes": extract_arts(wb),
        "mecenas": extract_patrons(wb),
        "sombra": extract_shadow_rules(wb),
        "nombres_por_cultura": extract_names_by_culture(wb),
    }
    
    wb.close()
    return data

if __name__ == "__main__":
    data = extract_all_data()
    
    # Save to JSON for inspection
    output_path = DATA_DIR / 'extracted_data.json'
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    
    print(f"Data extracted and saved to {output_path}")
    print(f"\nSummary:")
    print(f"  - Culturas: {len(data['culturas'])}")
    print(f"  - Trasfondos: {len(data['trasfondos'])}")
    print(f"  - Ocupaciones: {len(data['ocupaciones'])}")
    print(f"  - Virtudes: {len(data['virtudes'])}")
    print(f"  - Artes: {len(data['artes'])}")
    print(f"  - Mecenas: {len(data['mecenas'])}")
