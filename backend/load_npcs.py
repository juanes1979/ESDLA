"""
Script to load NPCs, enemies, and animals from Excel file into MongoDB
"""
import openpyxl
import asyncio
from motor.motor_asyncio import AsyncIOMotorClient
import os
from dotenv import load_dotenv
from datetime import datetime, timezone
import uuid

load_dotenv('.env')

def now_utc():
    return datetime.now(timezone.utc).isoformat()

async def load_npcs():
    client = AsyncIOMotorClient(os.environ['MONGO_URL'])
    db = client[os.environ['DB_NAME']]
    
    wb = openpyxl.load_workbook('/tmp/npcs.xlsx')
    ws = wb['Hoja1']
    
    category_map = {
        'MALIGNOS': 'malignos',
        'PNJ': 'pnj',
        'ANIMALES': 'animales',
        'ESPECIALES': 'especiales'
    }
    
    all_npcs = {
        'malignos': [],
        'pnj': [],
        'animales': [],
        'especiales': []
    }
    
    current_category = None
    
    for row in range(1, ws.max_row + 1):
        first_cell = ws.cell(row, 1).value
        
        # Check for category header
        if first_cell in category_map:
            current_category = category_map[first_cell]
            continue
        
        # Skip header rows
        if first_cell == 'Nombre' or not first_cell:
            continue
        
        # Extract NPC data
        if current_category:
            # Parse velocidad - can be number or string like "3 (correr)"
            velocidad_raw = ws.cell(row, 6).value
            if isinstance(velocidad_raw, str):
                # Extract number from string
                import re
                match = re.search(r'(\d+(?:\.\d+)?)', velocidad_raw)
                velocidad = float(match.group(1)) if match else 9
                velocidad_nota = velocidad_raw
            else:
                velocidad = velocidad_raw or 9
                velocidad_nota = None
            
            npc = {
                '_id': str(uuid.uuid4()),
                'nombre': first_cell,
                'categoria': current_category,
                'descripcion': ws.cell(row, 2).value or '',
                'tipo': ws.cell(row, 3).value or '',  # Can be race type or alignment
                'clase_armadura': ws.cell(row, 4).value or 10,
                'puntos_golpe': ws.cell(row, 5).value or 1,
                'velocidad': velocidad,
                'velocidad_nota': velocidad_nota,
                'atributos': {
                    'fuerza': ws.cell(row, 8).value or 10,
                    'destreza': ws.cell(row, 9).value or 10,
                    'constitucion': ws.cell(row, 10).value or 10,
                    'inteligencia': ws.cell(row, 11).value or 10,
                    'sabiduria': ws.cell(row, 12).value or 10,
                    'carisma': ws.cell(row, 13).value or 10
                },
                'sentidos': ws.cell(row, 14).value or '',
                'experiencia': ws.cell(row, 15).value or 0,
                'especial': ws.cell(row, 16).value or '',
                'acciones': ws.cell(row, 17).value or '',
                'created_at': now_utc(),
                'updated_at': now_utc()
            }
            
            all_npcs[current_category].append(npc)
    
    # Clear existing data and insert new
    await db.npcs.delete_many({})
    
    # Insert all NPCs
    all_records = []
    for category, npcs in all_npcs.items():
        all_records.extend(npcs)
    
    if all_records:
        await db.npcs.insert_many(all_records)
    
    print(f"Loaded {len(all_records)} NPCs:")
    for category, npcs in all_npcs.items():
        print(f"  - {category}: {len(npcs)}")
    
    return all_npcs

if __name__ == "__main__":
    asyncio.run(load_npcs())
