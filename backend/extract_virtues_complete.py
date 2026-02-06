"""
Script to extract COMPLETE virtue data from utumno.xlsm
Structure:
- Row 2: Virtue name
- Row 3: Description
- Row 4: Traits to note on character sheet
- Rows 5-10: Direct stat increases (x = increase this stat by 1)
- Rows 12-17: Choose ONE stat to increase by 1 (x = available option)
- Rows 19-24: Choose ONE saving throw proficiency (x = available option)
- Row 26: Hit points bonus
- Row 27: Community points bonus
- Row 28: Armor class bonus
- Rows 30-48: Choose ONE skill proficiency (x = available option)
- Rows 50-68: Choose ONE tool proficiency (x = available option)
"""
import openpyxl
import asyncio
from motor.motor_asyncio import AsyncIOMotorClient
import os
import uuid
from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

def safe_str(val, max_len=1000):
    if val is None:
        return ""
    return str(val)[:max_len].strip()

def is_marked(val):
    """Check if cell has an 'x' mark"""
    if val is None:
        return False
    return str(val).strip().lower() == 'x'

def safe_int(val):
    if val is None:
        return None
    try:
        return int(val)
    except:
        return None

# Row definitions for virtue structure
STAT_NAMES = ['FUERZA', 'DESTREZA', 'CONSTITUCIÓN', 'INTELIGENCIA', 'SABIDURÍA', 'CARISMA']
STAT_ROWS_DIRECT = list(range(5, 11))  # Rows 5-10 for direct increases
STAT_ROWS_CHOICE = list(range(12, 18))  # Rows 12-17 for choice increases
SAVE_ROWS = list(range(19, 25))  # Rows 19-24 for saving throw proficiencies
SKILL_ROWS = list(range(30, 49))  # Rows 30-48 for skill proficiencies
TOOL_ROWS = list(range(50, 69))  # Rows 50-68 for tool proficiencies

async def extract_virtues():
    wb = openpyxl.load_workbook('/app/data/utumno.xlsm', data_only=True)
    ws = wb['Virtudes']
    
    # Get skill names from column B
    skill_names = {}
    for row in SKILL_ROWS:
        skill_names[row] = safe_str(ws.cell(row=row, column=2).value)
    
    # Get tool names from column B
    tool_names = {}
    for row in TOOL_ROWS:
        tool_names[row] = safe_str(ws.cell(row=row, column=2).value)
    
    virtues = []
    
    # Iterate through all virtue columns (3 onwards)
    for col in range(3, ws.max_column + 1):
        virtue_name = safe_str(ws.cell(row=2, column=col).value)
        if not virtue_name or virtue_name == 'VIRTUD':
            continue
        
        # Get culture from row 1
        cultura = safe_str(ws.cell(row=1, column=col).value)
        if not cultura:
            # Look backwards for culture header
            for prev_col in range(col-1, 0, -1):
                prev_cult = safe_str(ws.cell(row=1, column=prev_col).value)
                if prev_cult:
                    cultura = prev_cult
                    break
        
        # Basic info
        descripcion = safe_str(ws.cell(row=3, column=col).value)
        rasgos = safe_str(ws.cell(row=4, column=col).value)
        
        # Direct stat increases (rows 5-10)
        aumenta_fuerza = is_marked(ws.cell(row=5, column=col).value)
        aumenta_destreza = is_marked(ws.cell(row=6, column=col).value)
        aumenta_constitucion = is_marked(ws.cell(row=7, column=col).value)
        aumenta_inteligencia = is_marked(ws.cell(row=8, column=col).value)
        aumenta_sabiduria = is_marked(ws.cell(row=9, column=col).value)
        aumenta_carisma = is_marked(ws.cell(row=10, column=col).value)
        
        # Stat choices (rows 12-17) - choose ONE
        stat_choices = []
        for i, row in enumerate(STAT_ROWS_CHOICE):
            if is_marked(ws.cell(row=row, column=col).value):
                stat_choices.append(STAT_NAMES[i])
        
        # Saving throw proficiency choices (rows 19-24)
        salvacion_choices = []
        for i, row in enumerate(SAVE_ROWS):
            if is_marked(ws.cell(row=row, column=col).value):
                salvacion_choices.append(STAT_NAMES[i])
        
        # Bonuses (rows 26-28)
        bonus_pg = safe_int(ws.cell(row=26, column=col).value)
        bonus_comunidad = safe_int(ws.cell(row=27, column=col).value)
        bonus_ca = safe_int(ws.cell(row=28, column=col).value)
        
        # Skill proficiency choices (rows 30-48)
        habilidad_choices = []
        for row in SKILL_ROWS:
            if is_marked(ws.cell(row=row, column=col).value):
                skill = skill_names.get(row)
                if skill:
                    habilidad_choices.append(skill)
        
        # Tool proficiency choices (rows 50-68)
        herramienta_choices = []
        for row in TOOL_ROWS:
            if is_marked(ws.cell(row=row, column=col).value):
                tool = tool_names.get(row)
                if tool:
                    herramienta_choices.append(tool)
        
        virtue = {
            "_id": str(uuid.uuid4()),
            "nombre": virtue_name,
            "cultura": cultura if cultura != "COMUNES" else None,
            "es_comun": cultura == "COMUNES",
            "descripcion": descripcion,
            "rasgos": rasgos,
            # Direct stat increases
            "aumenta_fuerza": aumenta_fuerza,
            "aumenta_destreza": aumenta_destreza,
            "aumenta_constitucion": aumenta_constitucion,
            "aumenta_inteligencia": aumenta_inteligencia,
            "aumenta_sabiduria": aumenta_sabiduria,
            "aumenta_carisma": aumenta_carisma,
            # Choice-based stat increase (choose ONE from list)
            "elegir_caracteristica": stat_choices if stat_choices else None,
            # Saving throw proficiency choices (choose ONE from list)
            "elegir_salvacion": salvacion_choices if salvacion_choices else None,
            # Bonuses
            "bonus_puntos_golpe": bonus_pg,
            "bonus_comunidad": bonus_comunidad,
            "bonus_ca": bonus_ca,
            # Skill proficiency choices (choose ONE from list)
            "elegir_habilidad": habilidad_choices if habilidad_choices else None,
            # Tool proficiency choices (choose ONE from list)
            "elegir_herramienta": herramienta_choices if herramienta_choices else None,
        }
        
        virtues.append(virtue)
    
    # Clear existing and insert new
    await db.virtues.delete_many({})
    if virtues:
        await db.virtues.insert_many(virtues)
    
    print(f"Extracted {len(virtues)} virtues")
    
    # Summary
    cultures = set(v['cultura'] for v in virtues if v['cultura'])
    print(f"Cultures: {len(cultures)}")
    
    with_stat_choice = sum(1 for v in virtues if v['elegir_caracteristica'])
    with_save_choice = sum(1 for v in virtues if v['elegir_salvacion'])
    with_skill_choice = sum(1 for v in virtues if v['elegir_habilidad'])
    with_tool_choice = sum(1 for v in virtues if v['elegir_herramienta'])
    
    print(f"With stat choice: {with_stat_choice}")
    print(f"With save choice: {with_save_choice}")
    print(f"With skill choice: {with_skill_choice}")
    print(f"With tool choice: {with_tool_choice}")
    
    # Sample output
    print("\n=== SAMPLE VIRTUES ===")
    for v in virtues[:3]:
        print(f"\n{v['nombre']} ({v['cultura'] or 'Común'}):")
        print(f"  Descripción: {v['descripcion'][:60]}...")
        print(f"  Rasgos: {v['rasgos'][:60] if v['rasgos'] else 'N/A'}...")
        if v['aumenta_fuerza'] or v['aumenta_destreza'] or v['aumenta_constitucion']:
            print(f"  Aumenta: FUE={v['aumenta_fuerza']}, DES={v['aumenta_destreza']}, CON={v['aumenta_constitucion']}")
        if v['elegir_caracteristica']:
            print(f"  Elegir característica: {v['elegir_caracteristica']}")
        if v['elegir_salvacion']:
            print(f"  Elegir salvación: {v['elegir_salvacion']}")
        if v['elegir_habilidad']:
            print(f"  Elegir habilidad: {v['elegir_habilidad']}")
        if v['elegir_herramienta']:
            print(f"  Elegir herramienta: {v['elegir_herramienta']}")
        if v['bonus_pg']:
            print(f"  +{v['bonus_pg']} PG")
        if v['bonus_ca']:
            print(f"  +{v['bonus_ca']} CA")

if __name__ == "__main__":
    asyncio.run(extract_virtues())
